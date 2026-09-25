"""Bicim donusturme: PDF'ten disariya, disaridan PDF'e.

Tasarim notu: her donusum "ne kadar sadik?" sorusuna dururst cevap verir.
Yuksek sadakat gerektiren yollar (DOCX -> PDF) kurulu Word'u kullanir; Word
yoksa metin tabanli yedege duser ve cagiran tarafa bunu bildirir.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile

import pymupdf

from .model import escape_html

# PyMuPDF'in dogrudan acip PDF'e cevirebildigi belge bicimleri.
NATIVE_DOC_EXT = {".epub", ".xps", ".oxps", ".cbz", ".fb2", ".mobi", ".svg"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp",
             ".pnm", ".pgm", ".ppm", ".jxr"}
TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".log"}
HTML_EXT = {".html", ".htm"}
OFFICE_EXT = {".docx", ".doc", ".rtf", ".odt", ".pptx", ".ppt", ".xlsx", ".xls"}


class ConversionError(Exception):
    pass


class Result:
    """Donusumun sonucu ve ne kadar sadik oldugu."""

    def __init__(self, path, note="", faithful=True):
        self.path = path
        self.note = note          # kullaniciya gosterilecek aciklama
        self.faithful = faithful  # bicim birebir korundu mu


# ---------------------------------------------------------------------------
# PDF -> disariya
# ---------------------------------------------------------------------------

def to_images(doc, folder: str, dpi: int = 150, fmt: str = "png",
              pages=None, stem: str = "sayfa") -> Result:
    zoom = dpi / 72.0
    ext = "jpg" if fmt.lower() in ("jpg", "jpeg") else "png"
    indices = list(pages) if pages is not None else range(doc.page_count)

    count = 0
    for i in indices:
        pix = doc.load_page(i).get_pixmap(matrix=pymupdf.Matrix(zoom, zoom),
                                          alpha=False)
        pix.save(os.path.join(folder, "%s-%03d.%s" % (stem, i + 1, ext)))
        count += 1
    return Result(folder, "%d sayfa %d DPI %s olarak kaydedildi."
                  % (count, dpi, ext.upper()))


def to_text(doc, path: str) -> Result:
    parts = []
    for i in range(doc.page_count):
        parts.append("--- Sayfa %d ---\n" % (i + 1))
        parts.append(doc.load_page(i).get_text("text"))
        parts.append("\n")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("".join(parts))
    return Result(path, "Düz metin çıkarıldı; sayfa düzeni ve görseller "
                        "korunmaz.", faithful=False)


def to_html(doc, path: str) -> Result:
    """Sayfalari konumlandirilmis HTML olarak yaz (gorunum buyuk olcude korunur)."""
    govde = []
    for i in range(doc.page_count):
        govde.append('<div class="sayfa">')
        govde.append(doc.load_page(i).get_text("html"))
        govde.append("</div>")

    html = (
        "<!DOCTYPE html><html lang='tr'><head><meta charset='utf-8'>"
        "<title>%s</title><style>"
        "body{background:#525659;margin:0;padding:20px;}"
        ".sayfa{background:#fff;margin:0 auto 20px;box-shadow:0 2px 8px "
        "rgba(0,0,0,.4);position:relative;overflow:hidden;}"
        "</style></head><body>%s</body></html>"
        % (escape_html(os.path.basename(path)), "".join(govde))
    )
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(html)
    return Result(path, "HTML olarak kaydedildi; fontlar tarayıcıdakilerle "
                        "değiştirilir.", faithful=False)


def to_svg(doc, folder: str, pages=None, stem: str = "sayfa") -> Result:
    indices = list(pages) if pages is not None else range(doc.page_count)
    for i in indices:
        svg = doc.load_page(i).get_svg_image()
        with open(os.path.join(folder, "%s-%03d.svg" % (stem, i + 1)),
                  "w", encoding="utf-8") as fh:
            fh.write(svg)
    return Result(folder, "%d sayfa SVG (vektör) olarak kaydedildi."
                  % len(list(indices)))


def to_docx(doc, path: str) -> Result:
    """Metni ve resimleri DOCX'e aktar.

    PDF'in yerlesimi Word'e birebir tasinamaz; paragraf akisi ve resimler
    korunur, sutun/kutu duzeni korunmaz.
    """
    try:
        import docx
        from docx.shared import Pt, RGBColor
    except ImportError:
        raise ConversionError(
            "DOCX çıktısı için python-docx gerekli:\n"
            "    .venv\\Scripts\\python.exe -m pip install python-docx")

    out = docx.Document()
    tmpdir = tempfile.mkdtemp(prefix="pdfstudio-docx-")

    try:
        for i in range(doc.page_count):
            page = doc.load_page(i)
            data = page.get_text("dict")

            for block in data.get("blocks", []):
                if block.get("type") == 1:          # resim blogu
                    _docx_image(out, block, tmpdir, docx)
                    continue

                for line in block.get("lines", []):
                    spans = line.get("spans", [])
                    text = "".join(s["text"] for s in spans).strip()
                    if not text:
                        continue
                    para = out.add_paragraph()
                    lead = max(spans, key=lambda s: len(s.get("text", "")))
                    run = para.add_run(text)
                    run.font.size = Pt(max(6.0, float(lead.get("size", 11))))
                    c = int(lead.get("color", 0))
                    run.font.color.rgb = RGBColor((c >> 16) & 255,
                                                  (c >> 8) & 255, c & 255)
                    flags = int(lead.get("flags", 0))
                    run.bold = bool(flags & 2 ** 4)
                    run.italic = bool(flags & 2 ** 1)

            if i < doc.page_count - 1:
                out.add_page_break()

        out.save(path)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

    return Result(path, "Metin ve resimler Word'e aktarıldı; sütun/kutu "
                        "yerleşimi korunmaz.", faithful=False)


def _docx_image(out, block, tmpdir, docx_mod) -> None:
    """DOCX'e tek bir resim blogu ekle; basarisiz olursa sessizce atla."""
    from docx.shared import Pt

    data = block.get("image")
    if not data:
        return
    ext = block.get("ext", "png")
    name = os.path.join(tmpdir, "img-%d.%s" % (id(block) % 10 ** 8, ext))
    try:
        with open(name, "wb") as fh:
            fh.write(data)
        bbox = pymupdf.Rect(block["bbox"])
        out.add_picture(name, width=Pt(min(468.0, bbox.width)))
    except Exception:
        # Desteklenmeyen resim bicimi belgenin tamamini bozmasin.
        pass


# ---------------------------------------------------------------------------
# disaridan -> PDF
# ---------------------------------------------------------------------------

def word_executable():
    """Kurulu Word varsa yolunu ver."""
    candidates = [
        r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE",
        r"C:\Program Files (x86)\Microsoft Office\root\Office16\WINWORD.EXE",
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    return shutil.which("winword")


def libreoffice_executable():
    for path in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(path):
            return path
    return shutil.which("soffice")


def images_to_pdf(paths, out_path: str) -> Result:
    """Her resmi bir sayfa yaparak PDF uret."""
    if not paths:
        raise ConversionError("Dönüştürülecek resim seçilmedi.")

    out = pymupdf.open()
    eklenen = 0
    atlanan = []
    for path in paths:
        try:
            img = pymupdf.open(path)
            pdf_bytes = img.convert_to_pdf()
            img.close()
            page_doc = pymupdf.open("pdf", pdf_bytes)
            out.insert_pdf(page_doc)
            page_doc.close()
            eklenen += 1
        except Exception:
            atlanan.append(os.path.basename(path))

    if not eklenen:
        out.close()
        raise ConversionError("Hiçbir resim okunamadı.")

    out.save(out_path, garbage=4, deflate=True)
    out.close()

    note = "%d resim PDF'e dönüştürüldü." % eklenen
    if atlanan:
        note += " Okunamayan: %s" % ", ".join(atlanan[:5])
    return Result(out_path, note, faithful=not atlanan)


def text_to_pdf(path: str, out_path: str, size: float = 11.0) -> Result:
    """Duz metin / Markdown dosyasini sayfalara dok.

    Sayfalama Story motoruyla yapiliyor: insert_htmlbox sigmayan metni
    kucultup tek sayfaya tikistiriyor, Story ise gercekten yeni sayfaya
    tasiyor.
    """
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        content = fh.read()

    html = ("<html><body><pre>%s</pre></body></html>"
            % escape_html(content))
    css = ("pre {font-family: sans-serif; font-size: %.1fpx; "
           "line-height: 1.45; white-space: pre-wrap; margin: 0;}" % size)

    _story_to_pdf(html, css, out_path, os.path.dirname(os.path.abspath(path)))
    return Result(out_path, "Metin PDF'e döküldü.")


def _story_to_pdf(html: str, css, out_path: str, base_dir: str) -> int:
    """HTML'i Story ile sayfalara bolerek PDF yaz; sayfa sayisini dondur."""
    try:
        archive = pymupdf.Archive(base_dir) if base_dir else None
        story = pymupdf.Story(html=html, user_css=css, archive=archive)
    except Exception as exc:
        raise ConversionError("İçerik işlenemedi: %s" % exc)

    writer = pymupdf.DocumentWriter(out_path)
    rect = pymupdf.paper_rect("a4")
    area = rect + (56, 56, -56, -56)

    more = 1
    pages = 0
    while more and pages < 5000:
        device = writer.begin_page(rect)
        more, _filled = story.place(area)
        story.draw(device)
        writer.end_page()
        pages += 1
    writer.close()
    return pages


def html_to_pdf(path: str, out_path: str) -> Result:
    """HTML dosyasini PDF'e cevir (Story motoru)."""
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        html = fh.read()

    base = os.path.dirname(os.path.abspath(path))
    _story_to_pdf(html, None, out_path, base)
    return Result(out_path, "HTML PDF'e dönüştürüldü; harici CSS ve "
                            "betikler uygulanmaz.", faithful=False)


def native_to_pdf(path: str, out_path: str) -> Result:
    """EPUB / XPS / CBZ / FB2 / SVG -> PDF (PyMuPDF dogrudan aciyor)."""
    try:
        src = pymupdf.open(path)
    except Exception as exc:
        raise ConversionError("Dosya açılamadı: %s" % exc)

    try:
        if src.is_pdf:
            src.save(out_path, garbage=4, deflate=True)
        else:
            data = src.convert_to_pdf()
            out = pymupdf.open("pdf", data)
            out.save(out_path, garbage=4, deflate=True)
            out.close()
    finally:
        src.close()
    return Result(out_path, "PDF'e dönüştürüldü.")


def office_to_pdf(path: str, out_path: str) -> Result:
    """Office belgesi -> PDF.

    Once kurulu Word, sonra LibreOffice denenir (ikisi de bicimi korur).
    Hicbiri yoksa DOCX icin metin tabanli yedek kullanilir ve bu durum
    cagirana bildirilir.
    """
    word = word_executable()
    if word:
        try:
            return _via_word(path, out_path)
        except Exception:
            pass          # COM yoksa/basarisizsa digerlerini dene

    soffice = libreoffice_executable()
    if soffice:
        try:
            return _via_libreoffice(soffice, path, out_path)
        except Exception:
            pass

    if path.lower().endswith(".docx"):
        return _docx_text_fallback(path, out_path)

    raise ConversionError(
        "Bu biçim için Microsoft Word veya LibreOffice gerekiyor; "
        "ikisi de bulunamadı.")


def _via_word(path: str, out_path: str) -> Result:
    import win32com.client                     # pywin32
    import pythoncom

    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(os.path.abspath(path), ReadOnly=True,
                                  AddToRecentFiles=False)
        doc.SaveAs(os.path.abspath(out_path), FileFormat=17)   # 17 = PDF
        return Result(out_path, "Microsoft Word ile dönüştürüldü; biçim "
                                "birebir korundu.")
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def _via_libreoffice(soffice: str, path: str, out_path: str) -> Result:
    outdir = tempfile.mkdtemp(prefix="pdfstudio-lo-")
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf",
             "--outdir", outdir, os.path.abspath(path)],
            check=True, timeout=180,
            stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        produced = os.path.join(
            outdir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
        if not os.path.exists(produced):
            raise ConversionError("LibreOffice çıktı üretmedi.")
        shutil.move(produced, out_path)
    finally:
        shutil.rmtree(outdir, ignore_errors=True)
    return Result(out_path, "LibreOffice ile dönüştürüldü; biçim büyük "
                            "ölçüde korundu.")


def _docx_text_fallback(path: str, out_path: str) -> Result:
    """Word/LibreOffice yokken DOCX'in metnini PDF'e dok."""
    try:
        import docx
    except ImportError:
        raise ConversionError(
            "Word veya LibreOffice bulunamadı ve python-docx kurulu değil.")

    document = docx.Document(path)
    satirlar = [p.text for p in document.paragraphs]
    tmp = tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False,
                                      encoding="utf-8")
    try:
        tmp.write("\n".join(satirlar))
        tmp.close()
        text_to_pdf(tmp.name, out_path)
    finally:
        os.unlink(tmp.name)

    return Result(out_path,
                  "Word bulunamadı; yalnızca metin aktarıldı. Biçim, tablo "
                  "ve görseller korunmadı.", faithful=False)


def any_to_pdf(path: str, out_path: str) -> Result:
    """Uzantiya bakarak dogru donusturucuyu sec."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        return native_to_pdf(path, out_path)
    if ext in IMAGE_EXT:
        return images_to_pdf([path], out_path)
    if ext in TEXT_EXT:
        return text_to_pdf(path, out_path)
    if ext in HTML_EXT:
        return html_to_pdf(path, out_path)
    if ext in NATIVE_DOC_EXT:
        return native_to_pdf(path, out_path)
    if ext in OFFICE_EXT:
        return office_to_pdf(path, out_path)

    raise ConversionError("Desteklenmeyen biçim: %s" % (ext or "(uzantısız)"))


def import_filter() -> str:
    """Dosya secme diyalogu icin filtre metni."""
    def join(exts):
        return " ".join("*" + e for e in sorted(exts))

    hepsi = join(IMAGE_EXT | TEXT_EXT | HTML_EXT | NATIVE_DOC_EXT
                 | OFFICE_EXT | {".pdf"})
    return (
        "Dönüştürülebilir dosyalar (%s);;"
        "Resimler (%s);;"
        "Office belgeleri (%s);;"
        "Metin (%s);;"
        "Web (%s);;"
        "E-kitap ve diğer (%s);;"
        "Tüm dosyalar (*)"
        % (hepsi, join(IMAGE_EXT), join(OFFICE_EXT), join(TEXT_EXT),
           join(HTML_EXT), join(NATIVE_DOC_EXT))
    )
