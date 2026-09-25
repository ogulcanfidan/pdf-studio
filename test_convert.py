"""Tum donusum yollarini gercek dosyalarla dener."""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pymupdf  # noqa: E402

from pdfstudio import convert  # noqa: E402

TMP = tempfile.mkdtemp(prefix="pdfstudio-conv-")
fails = []


def check(name, fn):
    try:
        res = fn()
        if isinstance(res, convert.Result):
            var = os.path.exists(res.path)
            boyut = (os.path.getsize(res.path)
                     if var and os.path.isfile(res.path) else -1)
            if not var:
                raise AssertionError("cikti olusmadi: %s" % res.path)
            if os.path.isfile(res.path) and boyut < 50:
                raise AssertionError("cikti bos (%d bayt)" % boyut)
            print("  OK   %-26s %s" % (name, res.note))
        else:
            print("  OK   %s" % name)
    except Exception as exc:
        print("  FAIL %-26s %r" % (name, exc))
        fails.append(name)


# --- kaynak dosyalar ------------------------------------------------------
src_pdf = os.path.join(TMP, "kaynak.pdf")
doc = pymupdf.open()
f = pymupdf.Font(fontfile="C:/Windows/Fonts/arial.ttf")
for i in range(3):
    page = doc.new_page()
    w = pymupdf.TextWriter(page.rect)
    w.append(pymupdf.Point(60, 90), "Sayfa %d — Türkçe başlık ĞÜŞİÖÇ" % (i + 1),
             font=f, fontsize=16)
    w.write_text(page, color=(0.1, 0.2, 0.5))
    w = pymupdf.TextWriter(page.rect)
    w.append(pymupdf.Point(60, 130), "Gövde metni, ikinci satır.", font=f,
             fontsize=11)
    w.write_text(page)
doc.save(src_pdf)
doc.close()

png = os.path.join(TMP, "resim.png")
pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 300, 200))
pix.set_rect(pymupdf.IRect(0, 0, 300, 200), (40, 110, 160))
pix.save(png)

txt = os.path.join(TMP, "not.txt")
with open(txt, "w", encoding="utf-8") as fh:
    fh.write("Başlık satırı\n\n" + ("Türkçe içerik ğüşiöç. " * 40 + "\n") * 8)

htm = os.path.join(TMP, "sayfa.html")
with open(htm, "w", encoding="utf-8") as fh:
    fh.write("<html><body><h1>Başlık ĞÜŞİ</h1><p>Paragraf metni. "
             "<b>Kalın</b> ve <i>italik</i>.</p></body></html>")

print("=== PDF -> disariya ===")
doc = pymupdf.open(src_pdf)
imgdir = os.path.join(TMP, "resimler")
os.makedirs(imgdir)
check("PDF -> PNG", lambda: convert.to_images(doc, imgdir, dpi=120))
check("PDF -> JPEG", lambda: convert.to_images(doc, imgdir, dpi=96, fmt="jpeg"))
check("PDF -> metin", lambda: convert.to_text(doc, os.path.join(TMP, "c.txt")))
check("PDF -> HTML", lambda: convert.to_html(doc, os.path.join(TMP, "c.html")))
svgdir = os.path.join(TMP, "svg")
os.makedirs(svgdir)
check("PDF -> SVG", lambda: convert.to_svg(doc, svgdir))
check("PDF -> DOCX", lambda: convert.to_docx(doc, os.path.join(TMP, "c.docx")))
doc.close()

# DOCX icerigi gercekten tasinmis mi?
try:
    import docx
    d = docx.Document(os.path.join(TMP, "c.docx"))
    metin = "\n".join(p.text for p in d.paragraphs)
    ok = "Türkçe başlık" in metin.replace("\xa0", " ")
    print("  %s DOCX icerik dogrulamasi -> %r"
          % ("OK  " if ok else "FAIL", metin[:60]))
    if not ok:
        fails.append("DOCX icerik")
except Exception as exc:
    print("  FAIL DOCX icerik okunamadi: %r" % exc)
    fails.append("DOCX icerik")

print()
print("=== disaridan -> PDF ===")
check("resim -> PDF",
      lambda: convert.images_to_pdf([png, png], os.path.join(TMP, "r.pdf")))
check("metin -> PDF",
      lambda: convert.text_to_pdf(txt, os.path.join(TMP, "t.pdf")))
check("HTML -> PDF",
      lambda: convert.html_to_pdf(htm, os.path.join(TMP, "h.pdf")))

# DOCX uret, sonra PDF'e cevir (Word yolu)
docx_path = os.path.join(TMP, "belge.docx")
try:
    import docx as _d
    w = _d.Document()
    w.add_heading("Başlık ĞÜŞİÖÇ", level=1)
    w.add_paragraph("Birinci paragraf, Türkçe karakterler: ğüşiöç.")
    w.add_paragraph("İkinci paragraf.")
    t = w.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Hücre A"
    t.cell(1, 1).text = "Hücre D"
    w.save(docx_path)
    check("DOCX -> PDF",
          lambda: convert.office_to_pdf(docx_path, os.path.join(TMP, "d.pdf")))
except Exception as exc:
    print("  FAIL DOCX hazirlanamadi: %r" % exc)
    fails.append("DOCX hazirlik")

print()
print("=== uretilen PDF'ler gecerli mi ===")
for ad in ("r.pdf", "t.pdf", "h.pdf", "d.pdf"):
    yol = os.path.join(TMP, ad)
    if not os.path.exists(yol):
        continue
    try:
        d = pymupdf.open(yol)
        metin = "".join(d.load_page(i).get_text("text")
                        for i in range(d.page_count))
        print("  OK   %-8s %d sayfa, %d karakter metin"
              % (ad, d.page_count, len(metin.strip())))
        d.close()
    except Exception as exc:
        print("  FAIL %-8s %r" % (ad, exc))
        fails.append(ad)

print()
print("=== uzun metin gercekten sayfalaniyor mu (kuculterek degil) ===")
tp = os.path.join(TMP, "t.pdf")
d = pymupdf.open(tp)
sayfa_sayisi = d.page_count
puntolar = []
for i in range(d.page_count):
    for b in d.load_page(i).get_text("dict")["blocks"]:
        for ln in b.get("lines", []):
            for sp in ln.get("spans", []):
                puntolar.append(sp["size"])
d.close()
en_kucuk = min(puntolar) if puntolar else 0
print("  sayfa sayisi: %d, en kucuk punto: %.2f" % (sayfa_sayisi, en_kucuk))
if sayfa_sayisi < 2:
    print("  FAIL uzun metin tek sayfaya sikismis")
    fails.append("sayfalama")
else:
    print("  OK   birden fazla sayfaya bolundu")
if en_kucuk < 8:
    print("  FAIL yazi okunamayacak kadar kucultulmus (%.2f pt)" % en_kucuk)
    fails.append("punto kuculmesi")
else:
    print("  OK   punto korundu (%.2f pt)" % en_kucuk)

print()
print("=== DOCX -> PDF icerik kontrolu (Word yolu) ===")
dp = os.path.join(TMP, "d.pdf")
if os.path.exists(dp):
    d = pymupdf.open(dp)
    metin = d.load_page(0).get_text("text").replace("\xa0", " ")
    d.close()
    for aranan in ("Başlık", "Birinci paragraf", "Hücre A"):
        ok = aranan in metin
        print("  %s %r" % ("OK  " if ok else "FAIL", aranan))
        if not ok:
            fails.append("docx icerik: " + aranan)

print()
print("=== any_to_pdf yonlendirmesi ===")
for kaynak in (png, txt, htm, docx_path, src_pdf):
    ad = os.path.basename(kaynak)
    hedef = os.path.join(TMP, "any-" + ad + ".pdf")
    try:
        res = convert.any_to_pdf(kaynak, hedef)
        print("  OK   %-14s -> %s" % (ad, res.note[:50]))
    except Exception as exc:
        print("  FAIL %-14s %r" % (ad, exc))
        fails.append("any:" + ad)

try:
    convert.any_to_pdf(os.path.join(TMP, "yok.xyz"), os.path.join(TMP, "x.pdf"))
    print("  FAIL desteklenmeyen bicim hata vermedi")
    fails.append("desteklenmeyen")
except convert.ConversionError:
    print("  OK   desteklenmeyen biçim düzgün hata veriyor")

print()
print("=" * 60)
if fails:
    print("BASARISIZ (%d): %s" % (len(fails), ", ".join(fails)))
    sys.exit(1)
print("HEPSI GECTI")
